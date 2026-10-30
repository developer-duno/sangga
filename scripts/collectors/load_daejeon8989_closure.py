# -*- coding: utf-8 -*-
"""대전시 상권분석 서비스(8989.daejeon.go.kr) 행정동 월별 폐업 → DB 적재 (결정 0044 · 한 트랜잭션).

무엇을 넣나
-----------
사람이 대전시 사이트 데이터셋 페이지에서 한 달씩 받아 둔 xlsx 를 **공표값 그대로** 넣는다:
  `<받은 날 폴더>/adstrd_clsbiz_monthly/daejeon8989_adstrd_clsbiz_YYYYMM.xlsx`
  (받은 날 폴더 = `data/raw/daejeon_8989/<YYYYMMDD>` — 폴더 이름이 곧 받은 날이다)
머리줄 = `행정동 명 | 업종분류명(소분류) | 폐업 년월 | 폐업 개수 | (빈 칸)` — 맨 끝 빈 칸이 있다.
업종은 한국표준산업분류 **소분류 이름**(소진공 업종과 다르다).

넣는 표 셋 + 받은 날 한 줄 (전부 통째 교체)
-------------------------------------------
  · daejeon_dong_closure (달, 행정동, 업종, 폐업 수) — **가게 업종(씨앗 y 33종) · 2025-01 이후 달만**
    (👤 결정 0044 ②③ — 2024년까지는 대전시 집계 방식이 달라 보인다 · 최근 6개월은 함수가 창에서 뺀다)
  · daejeon_pnu_dong (땅 PNU → 행정동) · daejeon_bjd_dong (법정동 10자리 → 행정동)
    — 대전시 자료에는 행정동 **이름**만 있어 건물(PNU)에 붙일 다리가 필요하다. 소진공 대전 점포 csv 의
    `지번코드`(19자리 = PNU)·`행정동명` 으로 만든다:
      ⓐ 지번코드 → 그 땅 점포가 가장 많은 행정동(동점이면 그 땅은 ⓐ 에서 뺀다 — 고를 근거가 없다)
      ⓑ 법정동(PNU 앞 10자리) → 행정동 — 소진공 자료에서 그 법정동이 **행정동 하나로만** 이어질 때만
    함수는 ⓐ 먼저, 없으면 ⓑ, 둘 다 없으면 「동 못 찾음」.
  · daejeon_closure_release (받은 날) — 한 줄. 함수의 창(받은 날의 달 − 6개월)과 신선도 줄이 이 값을 읽는다.

관문 (하나라도 걸리면 DB 는 손도 안 댄다 · exit 1)
--------------------------------------------------
  ⒜ 파일 0개
  ⒝ 머리줄이 위 넷(+ 끝 빈 칸)과 다르다
  ⒞ **씨앗(scripts/seeds/daejeon8989_industries.csv)에 없는 업종 이름** — 대전시가 업종을 바꾸면 멈춘다
     (조용히 빠지거나 '가게 아님'으로 새지 않게 · 사람이 씨앗에 y/n 을 적은 뒤 다시 돌린다)
  ⒟ 땅 → 동 지도에 없는 행정동 이름(대전시가 동을 나누거나 합친 날)
  ⒠ 폐업 년월이 파일 이름의 달과 다르다
  ⒡ **달 파일이 빠졌다** — 2025-01 부터 「받은 날의 달 − 1」까지 매달 파일이 하나씩 있어야 한다(대전시는 폐업이 없는 달도
     머리줄만 든 파일을 낸다 — 2026-10-28 실측 149파일 = 149달). 받은 날이 오늘보다 뒤여도 멈춘다(폴더 이름 오타).
     파일 중 가장 늦은 달이 받은 날의 달보다 뒤여도 멈춘다(폴더 이름을 과거로 잘못 적은 날 — 같은 달은 허용).
     ⛔ 통째 교체라서 — 몇 달만 받아 넣으면 나머지 달이 DB 에서 사라진다.
  ⒢ **반대 방향 동** — 땅 → 동 지도에 쓰인 동 이름이 대전시 원본(아무 달·업종)에 하나라도 없으면 멈춘다 · 소진공 csv 에서
     동 **이름 하나가 행정동코드 둘 이상**에 걸려도 멈춘다(동이 나뉘거나 이름이 바뀐 날 — 사람이 씨앗·지도를 본다).
  (+ 개수가 1 이상 정수가 아니거나 같은 (달·동·업종) 줄이 둘이면 멈춘다 — 표의 열쇠·check 가 어차피 막는 것을 먼저)
  (+ 실적재는 **받은 날 앞으로만** — DB 의 받은 날보다 이른 폴더면 적재 SQL 이 멈춘다 · 같은 날이면 다시 넣기 허용)

쓰는 법 (프로젝트 루트에서)
---------------------------
    python scripts/collectors/load_daejeon8989_closure.py --raw-dir data/raw/daejeon_8989/20261028 --dry-run
    python scripts/collectors/load_daejeon8989_closure.py --raw-dir data/raw/daejeon_8989/20261028
    → python scripts/post_load.py   그리고   python scripts/post_load.py --check
  `--store-csv` 를 안 주면 data/raw/sangkwon_*/ 의 가장 최근 대전 점포 csv 를 쓴다.
  `--dry-run` 은 DB 를 열지 않는다(지도·행 수·창·관문 결과만 찍는다) · 파일도 안 쓴다.
"""

import argparse
import collections
import csv
import datetime
import decimal
import glob
import io
import os
import re
import sys

SCRIPTS_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

PROJECT_ROOT = os.path.dirname(SCRIPTS_DIR)
SEED_CSV = os.path.join(SCRIPTS_DIR, "seeds", "daejeon8989_industries.csv")
DEFAULT_STAGING_DIR = os.path.join(PROJECT_ROOT, "data", "staging", "daejeon8989")
STORE_CSV_GLOB = os.path.join(PROJECT_ROOT, "data", "raw", "sangkwon_*",
                              "소상공인시장진흥공단_상가(상권)정보_대전_*.csv")

MONTHLY_DIR = "adstrd_clsbiz_monthly"
RE_FILE = re.compile(r"^daejeon8989_adstrd_clsbiz_(\d{4}(?:0[1-9]|1[0-2]))\.xlsx$")
RE_RECEIVED = re.compile(r"^\d{8}$")
RE_PNU = re.compile(r"^\d{19}$")
HEADER = ("행정동 명", "업종분류명(소분류)", "폐업 년월", "폐업 개수")

# 👤 결정 0044 ③ — 2025-01 부터 · 받은 날의 달에서 6개월 뺀 달까지 · 세 달이 모두 든 분기만 · 최근 8분기.
#    ⛔ 같은 수를 DB 함수 get_daejeon_dong_closure 의 const 가 쓴다(시험이 둘을 대조한다).
START_YM = "202501"
LAG_MONTHS = 6
MAX_QUARTERS = 8
SUM_QUARTERS = 4

CLOSURE_COLUMNS = ("ym", "dong_nm", "induty_nm", "clsbiz_co")


class GateError(ValueError):
    """관문에 걸림 — DB 에 아무것도 안 쓴다."""


# ── 순수 함수 ────────────────────────────────────────────────────────────────


def read_seed(path=SEED_CSV):
    """씨앗 → {업종 이름: 가게 업종인가}. 머리·y/n·중복이 틀리면 멈춘다."""
    with io.open(path, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    if not rows or rows[0] != ["induty_nm", "store"]:
        raise GateError("씨앗 머리줄이 induty_nm,store 가 아닙니다: {}".format(rows[:1]))
    out = {}
    for i, r in enumerate(rows[1:], 2):
        if len(r) != 2 or not r[0].strip() or r[1] not in ("y", "n"):
            raise GateError("씨앗 {}번째 줄 모양이 틀렸습니다: {!r}".format(i, r))
        if r[0] in out:
            raise GateError("씨앗에 같은 업종이 두 번: {!r}".format(r[0]))
        out[r[0]] = r[1] == "y"
    return out


def cell(v):
    """엑셀 칸 → 앞뒤 빈칸 뺀 글. None 은 빈 글."""
    return "" if v is None else str(v).strip()


def check_header(row, where):
    """⒝ 머리줄 = 넷 + (끝에 빈 칸만). 다르면 멈춘다."""
    cells = [cell(v) for v in (row or ())]
    while cells and cells[-1] == "":
        cells.pop()
    if tuple(cells) != HEADER:
        raise GateError("{}: ⒝ 머리줄이 다릅니다 — {} (기대 {})".format(where, cells, list(HEADER)))


def to_count(value, where):
    """폐업 개수 '3' · 3 · '3.0' → 3. 1 이상 정수가 아니면 멈춘다."""
    try:
        d = decimal.Decimal(cell(value))
    except decimal.InvalidOperation:
        raise GateError("{}: 폐업 개수가 수가 아닙니다: {!r}".format(where, value))
    if d != d.to_integral_value() or d < 1:
        raise GateError("{}: 폐업 개수가 1 이상 정수가 아닙니다: {!r}".format(where, value))
    return int(d)


def parse_month(rows, file_ym, where):
    """한 달 파일의 줄(머리 포함 튜플 목록) → [(달, 동, 업종, 개수)]. ⒝·⒠·개수를 본다."""
    if not rows:
        raise GateError("{}: ⒝ 머리줄이 없습니다(빈 파일)".format(where))
    check_header(rows[0], where)
    out = []
    for i, r in enumerate(rows[1:], 2):
        cells = [cell(v) for v in r]
        if not any(cells):
            continue
        cells += [""] * (4 - len(cells))
        dong, induty, ym, cnt = cells[0], cells[1], cells[2], r[3] if len(r) > 3 else None
        at = "{}:{}".format(where, i)
        if not dong or not induty:
            raise GateError("{}: 행정동·업종 이름이 비었습니다: {!r}".format(at, cells[:4]))
        if ym != file_ym:
            raise GateError("{}: ⒠ 폐업 년월 {!r} ≠ 파일 이름의 달 {}".format(at, ym, file_ym))
        out.append((ym, dong, induty, to_count(cnt, at)))
    return out


def gate_problems(rows, seed, dongs):
    """⒞·⒟·중복 — 문제 문장 목록(비면 통과). rows = [(달, 동, 업종, 개수)]."""
    bad = []
    unknown = sorted({r[2] for r in rows} - set(seed))
    if unknown:
        bad.append("⒞ 씨앗에 없는 업종 {}개 {} — scripts/seeds/daejeon8989_industries.csv 에 사람이 y/n 을 "
                   "적은 뒤 다시 돌린다".format(len(unknown), unknown[:10]))
    lost = sorted({r[1] for r in rows} - set(dongs))
    if lost:
        bad.append("⒟ 땅 → 동 지도에 없는 행정동 {}개 {} — 대전시가 동을 바꿨거나 소진공 점포 csv 가 낡았다"
                   .format(len(lost), lost[:10]))
    keys = collections.Counter((r[0], r[1], r[2]) for r in rows)
    dup = [k for k, n in keys.items() if n > 1]
    if dup:
        bad.append("같은 (달·동·업종) 줄이 둘 이상 {}개(예 {})".format(len(dup), dup[0]))
    return bad


def month_gaps(file_yms, received):
    """⒡ 2025-01 부터 「받은 날의 달 − 1」까지 빠진 달 파일 목록(옛것부터). 받은 날이 그 범위를 못 만들면 빈 목록."""
    have, out = set(file_yms), []
    last = ym_add(received.strftime("%Y%m"), -1)
    ym = START_YM
    while ym <= last:
        if ym not in have:
            out.append(ym)
        ym = ym_add(ym, 1)
    return out


def file_problems(file_yms, received, today):
    """⒡ — 문제 문장 목록(비면 통과). today 는 시험이 바꿔 넣는 자리(본체는 오늘)."""
    bad = []
    if received > today:
        bad.append("⒡ 받은 날 {} 이 오늘 {} 보다 뒤입니다 — 폴더 이름을 확인하세요".format(received, today))
    latest, recv_ym = max(file_yms, default=None), received.strftime("%Y%m")
    if latest is not None and latest > recv_ym:
        bad.append("⒡ 파일 중 가장 늦은 달 {} 이 받은 날 {} 의 달보다 뒤입니다 — 폴더 이름(받은 날)을 확인하세요"
                   .format(latest, received))
    gaps = month_gaps(file_yms, received)
    if gaps:
        bad.append("⒡ 빠진 달 파일 {}개 {} — 2025-01 부터 지난달까지 **전부 다시** 받아야 한다(통째 교체)"
                   .format(len(gaps), gaps))
    return bad


def reverse_dong_problems(map_dong_names, raw_dong_names, store_rows):
    """⒢ — 지도에 쓰인 동이 원본에 없음 · 소진공 csv 에서 동 이름 하나가 행정동코드 둘 이상. store_rows = [(지번, 동, 동코드)]."""
    bad = []
    lost = sorted(set(map_dong_names) - set(raw_dong_names))
    if lost:
        bad.append("⒢ 땅 → 동 지도의 동 {}개가 대전시 원본에 없습니다 {} — 동이 나뉘거나 이름이 바뀌었을 수 있다"
                   .format(len(lost), lost[:10]))
    codes = collections.defaultdict(set)
    for _pnu, dong, code in store_rows:
        dong, code = cell(dong), cell(code)
        if dong and code:
            codes[dong].add(code)
    multi = sorted(d for d, c in codes.items() if len(c) > 1)
    if multi:
        bad.append("⒢ 소진공 csv 에서 행정동 이름 하나가 코드 둘 이상에 걸립니다 {} — 이름으로 이을 수 없다"
                   .format({d: sorted(codes[d]) for d in multi[:5]}))
    return bad


def month_totals(rows):
    """dry-run 보고 — 달별 합(넣는 줄 · 옛 → 새)."""
    out = collections.Counter()
    for ym, _d, _i, n in rows:
        out[ym] += n
    return dict(sorted(out.items()))


def rows_to_load(rows, seed):
    """DB 에 넣는 줄 — 가게 업종(y) · 2025-01 이후 달만(👤 결정 0044 ②③)."""
    return [r for r in rows if seed.get(r[2]) is True and r[0] >= START_YM]


def build_dong_map(store_rows):
    """[(지번코드, 행정동명)] → (ⓐ {pnu: 동}, ⓑ {법정동 10자리: 동}, 통계 dict).

    ⓐ 땅마다 점포가 가장 많은 동 — 동점이면 그 땅은 뺀다.
    ⓑ 법정동(PNU 앞 10자리)이 소진공 자료에서 행정동 **하나로만** 이어질 때만.
    19자리 숫자가 아닌 지번코드·빈 동 이름 줄은 버린다(통계 dropped).
    """
    by_pnu = collections.defaultdict(collections.Counter)
    by_bjd = collections.defaultdict(set)
    dropped = 0
    for pnu, dong in store_rows:
        pnu, dong = cell(pnu), cell(dong)
        if not RE_PNU.match(pnu) or not dong:
            dropped += 1
            continue
        by_pnu[pnu][dong] += 1
        by_bjd[pnu[:10]].add(dong)
    pnu_map, ties, multi = {}, 0, 0
    for pnu, c in by_pnu.items():
        if len(c) > 1:
            multi += 1
        top = c.most_common(2)
        if len(top) > 1 and top[0][1] == top[1][1]:
            ties += 1
            continue
        pnu_map[pnu] = top[0][0]
    bjd_map = {b: next(iter(s)) for b, s in by_bjd.items() if len(s) == 1}
    stats = {"pnu_all": len(by_pnu), "pnu_multi": multi, "pnu_ties": ties, "pnu_map": len(pnu_map),
             "bjd_all": len(by_bjd), "bjd_map": len(bjd_map), "dropped": dropped,
             "dongs": len({d for c in by_pnu.values() for d in c})}
    return pnu_map, bjd_map, stats


def map_dongs(pnu_map, bjd_map):
    return set(pnu_map.values()) | set(bjd_map.values())


def received_date(raw_dir):
    """받은 날 = 받은 날 폴더 이름(YYYYMMDD). 꼴이 틀리거나 없는 날이면 멈춘다."""
    name = os.path.basename(os.path.normpath(raw_dir))
    if not RE_RECEIVED.match(name):
        raise GateError("받은 날 폴더 이름이 YYYYMMDD 가 아닙니다: {!r}".format(name))
    try:
        return datetime.datetime.strptime(name, "%Y%m%d").date()
    except ValueError:
        raise GateError("받은 날 폴더 이름이 없는 날입니다: {!r}".format(name))


def ym_add(ym, months):
    y, m = int(ym[:4]), int(ym[4:])
    n = y * 12 + (m - 1) + months
    return "{:04d}{:02d}".format(n // 12, n % 12 + 1)


def window_quarters(received):
    """받은 날 → 창의 분기 코드('YYYYQ' · 옛 → 새). M = 받은 날의 달 − 6개월.

    2025-01 부터 M 까지, **세 달이 모두 M 이하인 분기만** · 그중 최근 8분기.
    20261028 → M 202604 → ['20251', '20252', '20253', '20254', '20261'].
    """
    last = ym_add(received.strftime("%Y%m"), -LAG_MONTHS)
    out, ym = [], START_YM
    while ym_add(ym, 2) <= last:
        out.append("{}{}".format(ym[:4], (int(ym[4:]) + 2) // 3))
        ym = ym_add(ym, 3)
    return out[-MAX_QUARTERS:]


def quarter_of(ym):
    return "{}{}".format(ym[:4], (int(ym[4:]) + 2) // 3)


def dong_summary(rows, dong, quarters):
    """한 동의 창 분기별 합 · 최근 4분기 합 · 업종별(최근 4분기 · 많은 순) — dry-run 대조·시험용(함수와 같은 셈)."""
    per_q = {q: 0 for q in quarters}
    recent = set(quarters[-SUM_QUARTERS:])
    ind = collections.Counter()
    for ym, d, induty, n in rows:
        if d != dong:
            continue
        q = quarter_of(ym)
        if q in per_q:
            per_q[q] += n
        if q in recent:
            ind[induty] += n
    top = sorted(ind.items(), key=lambda kv: (-kv[1], kv[0]))
    return [per_q[q] for q in quarters], sum(ind.values()), top


# ── 파일 읽기 ────────────────────────────────────────────────────────────────


def read_xlsx_rows(path):
    """xlsx 첫 시트의 줄 → 튜플 목록. openpyxl 은 **여기서만** 쓴다(CI 에는 없다 — 함수 안 import)."""
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        raise GateError("openpyxl 이 필요합니다 (xlsx 를 읽는 유일한 곳):\n  python -m pip install openpyxl")
    import warnings  # noqa: PLC0415
    with warnings.catch_warnings():
        # 대전시 파일은 "Workbook contains no default style" 경고를 낸다(무해).
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
        finally:
            wb.close()


def list_month_files(raw_dir):
    """[(달, 경로)] — 이름 꼴이 맞는 파일만 · 달 차례."""
    d = os.path.join(raw_dir, MONTHLY_DIR)
    out = []
    if os.path.isdir(d):
        for n in sorted(os.listdir(d)):
            m = RE_FILE.match(n)
            if m:
                out.append((m.group(1), os.path.join(d, n)))
    return out


def collect(raw_dir, reader=None):
    """⒜ 파일 0개면 멈춘다. → (줄 전부, 파일 수). reader 는 시험이 엑셀 없이 줄을 넣을 때만 바꾼다."""
    files = list_month_files(raw_dir)
    if not files:
        raise GateError("⒜ 파일이 0개입니다: {}".format(os.path.join(raw_dir, MONTHLY_DIR)))
    reader = reader or read_xlsx_rows
    rows = []
    for ym, path in files:
        rows += parse_month(reader(path), ym, os.path.basename(path))
    return rows, len(files)


def default_store_csv():
    found = sorted(glob.glob(STORE_CSV_GLOB))
    return found[-1] if found else None


def read_store_csv(path):
    """소진공 대전 점포 csv(utf-8-sig) → [(지번코드, 행정동명, 행정동코드)]."""
    with io.open(path, encoding="utf-8-sig", newline="") as f:
        r = csv.DictReader(f)
        missing = [c for c in ("지번코드", "행정동명", "행정동코드") if c not in (r.fieldnames or [])]
        if missing:
            raise GateError("{}: 칸이 없습니다 {}".format(os.path.basename(path), missing))
        return [(row["지번코드"], row["행정동명"], row["행정동코드"]) for row in r]


# ── SQL ──────────────────────────────────────────────────────────────────────


def copy_path_literal(path):
    return path.replace("\\", "/").replace("'", "''")


def write_csv(rows, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with io.open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerows(rows)
    return len(rows)


def build_sql(paths, counts, received):
    """적재 SQL — 한 트랜잭션. 표 셋 통째 교체 + 받은 날 한 줄. 어긋나면 raise(psql 은 select 결과와 무관하게 0)."""
    return """begin;
set local statement_timeout = '600s';

create temp table stage_closure (like daejeon_dong_closure including defaults) on commit drop;
create temp table stage_pnu (like daejeon_pnu_dong including defaults) on commit drop;
create temp table stage_bjd (like daejeon_bjd_dong including defaults) on commit drop;

\\copy stage_closure (ym, dong_nm, induty_nm, clsbiz_co) from '{closure}' with (format csv, encoding 'UTF8')
\\copy stage_pnu (pnu, dong_nm) from '{pnu}' with (format csv, encoding 'UTF8')
\\copy stage_bjd (bjd_cd, dong_nm) from '{bjd}' with (format csv, encoding 'UTF8')

-- ── 행수 대조 — CSV 가 잘려도 \\copy 는 에러 없이 거기까지만 넣는다 ─────────────
do $$
begin
  if (select count(*) from stage_closure) <> {n_closure}
     or (select count(*) from stage_pnu) <> {n_pnu}
     or (select count(*) from stage_bjd) <> {n_bjd} then
    raise exception '올린 행 수가 원본에서 읽은 수(폐업 {n_closure} · 땅 {n_pnu} · 법정동 {n_bjd})와 다릅니다 — 통째로 되돌립니다';
  end if;
end $$;

-- 받은 날 앞으로만 — 지금 DB 의 받은 날보다 이른 판을 넣으면 최근 달이 사라진다(같은 날이면 다시 넣기 허용).
do $$
begin
  if exists (select 1 from daejeon_closure_release where received_on > date '{received}') then
    raise exception '지금 DB 의 받은 날이 {received} 보다 뒤입니다 — 옛 판으로 되돌리지 않습니다(통째로 되돌립니다)';
  end if;
end $$;

-- 통째 교체 — 대전시가 지난 달을 고쳐 다시 내기도 하니 받은 판 전부로 바꾼다.
delete from daejeon_dong_closure;
delete from daejeon_pnu_dong;
delete from daejeon_bjd_dong;
insert into daejeon_dong_closure (ym, dong_nm, induty_nm, clsbiz_co)
  select ym, dong_nm, induty_nm, clsbiz_co from stage_closure;
insert into daejeon_pnu_dong (pnu, dong_nm) select pnu, dong_nm from stage_pnu;
insert into daejeon_bjd_dong (bjd_cd, dong_nm) select bjd_cd, dong_nm from stage_bjd;

insert into daejeon_closure_release (id, received_on, loaded_at)
values (1, date '{received}', now())
on conflict (id) do update set received_on = excluded.received_on, loaded_at = excluded.loaded_at;

do $$
begin
  if (select count(*) from daejeon_dong_closure) <> {n_closure}
     or (select count(*) from daejeon_pnu_dong) <> {n_pnu}
     or (select count(*) from daejeon_bjd_dong) <> {n_bjd} then
    raise exception '표에 든 행 수가 기대와 다릅니다 — 통째로 되돌립니다';
  end if;
end $$;

select (select count(*) from daejeon_dong_closure) as closure_rows,
       (select sum(clsbiz_co) from daejeon_dong_closure) as closure_sum,
       (select count(*) from daejeon_pnu_dong) as pnu_rows,
       (select count(*) from daejeon_bjd_dong) as bjd_rows,
       (select received_on from daejeon_closure_release) as received_on;

commit;
""".format(closure=copy_path_literal(paths["closure"]), pnu=copy_path_literal(paths["pnu"]),
           bjd=copy_path_literal(paths["bjd"]), n_closure=counts["closure"], n_pnu=counts["pnu"],
           n_bjd=counts["bjd"], received=received.isoformat())


# ── 메인 ─────────────────────────────────────────────────────────────────────


def main(argv=None):
    try:
        if sys.stdout.isatty():
            sys.stdout.reconfigure(errors="replace")
        else:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    p = argparse.ArgumentParser(description="대전시 행정동 월별 폐업 → daejeon_dong_closure 적재",
                                allow_abbrev=False)
    p.add_argument("--raw-dir", required=True, help="받은 날 폴더(data/raw/daejeon_8989/YYYYMMDD)")
    p.add_argument("--store-csv", default=None, help="소진공 대전 점포 csv(기본 = 가장 최근 분기)")
    p.add_argument("--staging-dir", default=DEFAULT_STAGING_DIR)
    p.add_argument("--dry-run", action="store_true", help="지도·행 수·창·관문만 — DB·파일 쓰기 0")
    a = p.parse_args(argv)

    try:
        received = received_date(a.raw_dir)
        seed = read_seed()
        store_csv = a.store_csv or default_store_csv()
        if not store_csv or not os.path.exists(store_csv):
            raise GateError("소진공 대전 점포 csv 가 없습니다: {}".format(store_csv or STORE_CSV_GLOB))
        store_rows = read_store_csv(store_csv)
        pnu_map, bjd_map, stats = build_dong_map([(p, d) for p, d, _c in store_rows])
        rows, n_files = collect(a.raw_dir)
        file_yms = [ym for ym, _p in list_month_files(a.raw_dir)]
    except (GateError, OSError) as e:
        print("실패 — DB 에는 아무것도 쓰지 않았습니다: {}".format(e), file=sys.stderr)
        return 1

    load = rows_to_load(rows, seed)
    quarters = window_quarters(received)
    print("=" * 74)
    print("대전시 행정동 폐업 → daejeon_dong_closure {}".format("(dry-run)" if a.dry_run else ""))
    print("=" * 74)
    print("  받은 날 {} · 파일 {}개 · 읽은 줄 {:,} · 폐업 {:,}건".format(
        received.isoformat(), n_files, len(rows), sum(r[3] for r in rows)))
    print("  씨앗 업종 {} (가게 y {} · n {}) · 원본에 나온 업종 {}".format(
        len(seed), sum(seed.values()), len(seed) - sum(seed.values()), len({r[2] for r in rows})))
    print("  넣을 줄(가게 업종 · {} 이후) {:,} · 폐업 {:,}건 · 달 {} · 동 {}".format(
        START_YM, len(load), sum(r[3] for r in load), len({r[0] for r in load}), len({r[1] for r in load})))
    print("  땅 → 동 지도 ({})".format(os.path.basename(store_csv)))
    print("    ⓐ 지번코드 {:,} · 동이 둘 이상인 땅 {} · 그중 동점(뺌) {} · 지도 {:,}".format(
        stats["pnu_all"], stats["pnu_multi"], stats["pnu_ties"], stats["pnu_map"]))
    print("    ⓑ 법정동 {} · 행정동 하나뿐 {} · 행정동 {} · 버린 줄 {}".format(
        stats["bjd_all"], stats["bjd_map"], stats["dongs"], stats["dropped"]))
    print("  창(받은 날 − {}개월 · 세 달 다 든 분기 · 최근 {}분기) {} · 업종 표 = 최근 {}분기 {}".format(
        LAG_MONTHS, MAX_QUARTERS, quarters, SUM_QUARTERS, quarters[-SUM_QUARTERS:]))
    if quarters:
        sums = [0] * len(quarters)
        for d in sorted({r[1] for r in load}):
            per_q, _, _ = dong_summary(load, d, quarters)
            sums = [x + y for x, y in zip(sums, per_q)]
        print("    대전 전체 분기별 폐업(가게 업종) {}".format(dict(zip(quarters, sums))))

    print("  달별 폐업(가게 업종 · {} 이후) {}".format(START_YM, month_totals(load)))

    problems = (gate_problems(rows, seed, map_dongs(pnu_map, bjd_map))
                + file_problems(file_yms, received, datetime.date.today())
                + reverse_dong_problems(map_dongs(pnu_map, bjd_map), {r[1] for r in rows}, store_rows))
    if problems:
        print()
        print("관문에 걸렸습니다 — DB 에는 아무것도 쓰지 않았습니다:", file=sys.stderr)
        for m in problems:
            print("  " + m, file=sys.stderr)
        return 1
    print("  관문 통과 (⒜ · ⒝ · ⒞ · ⒟ · ⒠ · ⒡ · ⒢ · 중복)")

    if a.dry_run:
        print()
        print("--dry-run — DB·파일에 아무것도 쓰지 않았습니다.")
        return 0

    paths = {k: os.path.join(a.staging_dir, k + ".csv") for k in ("closure", "pnu", "bjd")}
    counts = {
        "closure": write_csv(sorted(load), paths["closure"]),
        "pnu": write_csv(sorted(pnu_map.items()), paths["pnu"]),
        "bjd": write_csv(sorted(bjd_map.items()), paths["bjd"]),
    }
    sql = build_sql(paths, counts, received)

    import dbx  # noqa: PLC0415  (dry-run 은 DB 설정 없이도 돌아야 한다)
    rc = dbx.run_sql(sql)
    if rc != 0:
        print("적재 실패 (psql 종료코드 {}). 트랜잭션이라 아무것도 안 들어갔습니다.".format(rc), file=sys.stderr)
        return rc
    print()
    print("  적재 완료 ({}).".format(datetime.datetime.now().strftime("%Y-%m-%d %H:%M")))
    print("  다음: python scripts/post_load.py   그리고   python scripts/post_load.py --check")
    return 0


if __name__ == "__main__":
    sys.exit(main())
